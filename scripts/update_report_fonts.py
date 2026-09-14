from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'RAG-Anything已实现功能参数介绍报告.docx'
OUT = ROOT / '多模态知识库参数报告.docx'


def set_rfonts(rpr, *, east='宋体', latin='Times New Roman'):
    if rpr is None:
        return
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), east)
    rfonts.set(qn('w:ascii'), latin)
    rfonts.set(qn('w:hAnsi'), latin)
    rfonts.set(qn('w:cs'), east)


def set_style_font(style, east='宋体'):
    style.font.name = 'Times New Roman'
    style._element.get_or_add_rPr()
    set_rfonts(style._element.rPr, east=east)


latin_re = re.compile(r'[A-Za-z0-9]')


def chunks(text):
    if not text:
        return []
    out = []
    start = 0
    current = bool(latin_re.match(text[0]))
    for i, ch in enumerate(text[1:], 1):
        flag = bool(latin_re.match(ch))
        if flag != current:
            out.append((text[start:i], current))
            start, current = i, flag
    out.append((text[start:], current))
    return out


def format_paragraph(paragraph):
    is_h1 = paragraph.style and paragraph.style.name.lower() in {'heading 1', '标题 1'}
    east = '黑体' if is_h1 else '宋体'
    for run in list(paragraph.runs):
        text = run.text
        if not text:
            run.font.name = 'Times New Roman'
            run._element.get_or_add_rPr()
            set_rfonts(run._element.rPr, east=east)
            continue
        parts = chunks(text)
        if len(parts) == 1:
            run.font.name = 'Times New Roman'
            run._element.get_or_add_rPr()
            set_rfonts(run._element.rPr, east=east)
            # East Asian glyphs use the requested Chinese face; Latin/digits use Times.
            run.font._element.rPr.rFonts.set(qn('w:eastAsia'), east)
            continue
        old = run._element
        parent = old.getparent()
        idx = parent.index(old)
        for part, is_latin in parts:
            new = OxmlElement('w:r')
            if old.rPr is not None:
                new.append(deepcopy(old.rPr))
            set_rfonts(new.get_or_add_rPr(), east=east, latin='Times New Roman')
            t = OxmlElement('w:t')
            if part[:1].isspace() or part[-1:].isspace():
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t.text = part
            new.append(t)
            parent.insert(idx, new)
            idx += 1
        parent.remove(old)


def walk_table(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                format_paragraph(p)
            for nested in cell.tables:
                walk_table(nested)


doc = Document(str(SRC))
first_paragraph = doc.paragraphs[0]
for _ in range(6):
    first_paragraph.insert_paragraph_before()

for style in doc.styles:
    if style.type == 1:  # paragraph styles
        set_style_font(style, east='黑体' if style.name.lower() in {'heading 1', '标题 1'} else '宋体')

for p in doc.paragraphs:
    format_paragraph(p)
for table in doc.tables:
    walk_table(table)
for section in doc.sections:
    for p in section.header.paragraphs + section.footer.paragraphs:
        format_paragraph(p)
    for table in section.header.tables + section.footer.tables:
        walk_table(table)

for p in reversed(doc.paragraphs):
    if p.text.startswith('交付结论：'):
        for run in list(p.runs):
            p._element.remove(run._element)
        break

doc.save(str(OUT))
print(OUT)
