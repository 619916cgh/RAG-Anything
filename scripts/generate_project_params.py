from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(__file__).resolve().parents[1] / "多模态知识库项目详细参数.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, size=10.5, bold=False, color="000000"):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text_paragraph(cell, text, *, bold=False, blue=False, first_line=True):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    if first_line:
        pf.first_line_indent = Pt(21)
    r = p.add_run(text)
    set_run_font(r, 10.5, bold=bold, color="0000FF" if blue else "000000")
    return p


def set_table_widths(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_module(table, index, name, paragraphs):
    cells = table.add_row().cells
    cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cells[2].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for cell in cells:
        set_cell_margins(cell)
        set_cell_border(cell)
    p = cells[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(index))
    set_run_font(r, 10.5)
    p = cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(name)
    set_run_font(r, 11, bold=True)
    # Remove the default empty paragraph in the detail cell.
    detail = cells[2]
    detail.paragraphs[0]._element.getparent().remove(detail.paragraphs[0]._element)
    for item in paragraphs:
        if isinstance(item, tuple):
            text, is_blue, is_bold = item
        else:
            text, is_blue, is_bold = item, False, False
        add_text_paragraph(detail, text, bold=is_bold, blue=is_blue, first_line=not is_blue)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.25)
section.header_distance = Inches(0.59)
section.footer_distance = Inches(0.69)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)
r = title.add_run("多模态知识库与智能体问答平台项目详细参数")
set_run_font(r, 20, bold=False)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
intro.paragraph_format.line_spacing = 1.15
intro.paragraph_format.space_after = Pt(8)
r = intro.add_run("本参数清单依据当前项目代码、配置、数据库迁移和项目总结整理，描述平台已形成的知识接入、处理、检索、问答、权限与运维能力；外部模型、浏览器、Worker、Docker 和生产环境验收需以对应受控证据为准。")
set_run_font(r, 10.5)

doc.add_paragraph()

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
set_table_widths(table, [0.55, 1.45, 4.25])
hdr = table.rows[0].cells
for cell, text in zip(hdr, ["序号", "模块名称", "详细参数"]):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "DDEBF7")
    set_cell_margins(cell, top=120, bottom=120)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run_font(r, 12, bold=True)

modules = [
    ("平台定位与总体能力", [
        ("平台面向教育和专业实训场景，提供从知识接入、文档解析、分块、向量与图谱索引，到混合检索、智能体问答、引用回溯和会话管理的完整链路。", False, False),
        ("项目优势：通过统一的平台闭环，把多模态资料处理、知识检索和智能体应用连接到同一套可审计流程中，便于教学内容沉淀与持续复用。", True, True),
        ("当前实现以仓库代码、配置和迁移为事实基线；真实外部 Provider、浏览器和生产环境结论不在本清单中替代。", False, False),
    ]),
    ("系统架构与运行链路", [
        ("系统采用 React/Vite 前端、FastAPI 路由层、业务服务、RAG 核心、PostgreSQL、向量与图谱存储、Redis 及文件工作区的分层结构。", False, False),
        ("上传接口创建带有 owner、知识库、设置快照和追踪标识的任务；Worker 读取任务快照完成解析、分块、Embedding、图谱写入和就绪状态更新。", False, False),
        ("项目优势：分层架构和任务快照机制使上传配置、处理状态和后续检索参数具有清晰的责任边界，适合多角色协作与问题追踪。", True, True),
    ]),
    ("知识接入与文档处理", [
        ("支持 txt、md、pdf、doc、docx、ppt、pptx、xls、xlsx、html、json、csv 以及常见图像格式；Office 文档可转换为 PDF 或 Markdown 后参与索引。", False, False),
        ("解析器目录支持 docling 等解析策略，并可按文件类型配置覆盖；图像和表格处理可由对应开关控制。", False, False),
        ("分块策略包括 fixed_size、recursive 和 sentence；常用默认目标为 800 token，重叠窗口默认约 100 token，具体值由任务快照固化。", False, False),
        ("项目优势：平台将文本、表格、图像、公式等内容统一纳入处理链路，并保留每次上传时的解析和分块设置，方便复现实验结果。", True, True),
        ("边界说明：上传扩展名白名单和全局错误脱敏仍需按当前安全整改与受控环境证据判断。", False, False),
    ]),
    ("知识库与内容管理", [
        ("支持知识库创建、列表、切换、元数据更新和删除，并按角色能力控制写入、管理和删除操作。", False, False),
        ("支持文档列表、摘要分页、下载、删除、批量删除、失败重试和检索修复；支持分块分页查看、搜索、编辑和标签更新。", False, False),
        ("图谱数据包含实体、关系、属性及关联 chunk，可通过接口和前端进行节点、边及实体查询。", False, False),
        ("项目优势：知识库、文档、分块和标签形成可维护的内容生命周期，既支持大批量资料管理，也保留针对单个文档的修复入口。", True, True),
    ]),
    ("混合检索与智能体问答", [
        ("检索支持 BM25、向量和图谱三路通道，并通过 Reciprocal Rank Fusion 融合结果；常用 RRF_K 默认值为 60。", False, False),
        ("支持 balanced、precise、broad 等预置检索模式，并可配置 retrieval_top_k、chunk_top_k、重排和通道裁剪。", False, False),
        ("Agent SSE 支持 accepted、检索状态、增量 token、引用和完成等事件；会话和关键状态写入 PostgreSQL。", False, False),
        ("项目优势：混合检索将关键词、语义向量和知识图谱互补结合，并通过流式回答与引用回溯提升问答的可解释性和交互效率。", True, True),
        ("真实 LLM、Embedding、VLM、ASR 的可达性、配额和延迟受部署环境影响，需要单独的真实账号和 Worker 验收。", False, False),
    ]),
    ("用户角色与安全控制", [
        ("系统采用五级 RBAC：super_admin、dept_admin、teacher、assistant、student；权限以 resource:action 形式表达。", False, False),
        ("受保护接口统一通过 require_permission() 和当前持久角色校验；前端隐藏仅是体验层，不能替代后端授权。", False, False),
        ("认证、角色、Token 撤销、审计和会话代际信息由 PostgreSQL 负责持久化；WebSocket 会复核撤销状态和 session generation。", False, False),
        ("项目优势：权限控制、Token 撤销、审计和会话复核在后端形成闭环，能够覆盖管理员、教师、助理和学生等不同使用边界。", True, True),
    ]),
    ("个人平台设置与配置治理", [
        ("设置按平台默认、个人设置、知识库资源设置及请求或任务覆盖进行合并，新上传任务固化合并后的解析、分块、模型和检索快照。", False, False),
        ("设置更新带 revision 和 fingerprint，并使用 expected_revision 检测并发冲突，冲突时返回 409。", False, False),
        ("项目优势：设置优先级和版本校验兼顾统一治理与个性化使用，既能保留个人默认值，也能避免并发修改互相覆盖。", True, True),
    ]),
    ("部署存储与运维参数", [
        ("部署由 app、postgres、redis、migrate、marker 和 nginx 等服务组成；应用通过内部网络运行，Nginx 提供外部入口。", False, False),
        ("健康检查区分 /api/live 和 /api/ready：前者检查进程存活，后者检查依赖可用性；Prometheus 指标和脱敏监控通道用于运行观察。", False, False),
        ("PostgreSQL 是认证、共享状态、迁移历史和业务仓库的权威后端；生产迁移前必须先完成可恢复备份确认并设置 MIGRATION_BACKUP_ACKNOWLEDGED=true。", False, False),
        ("项目优势：运行状态、迁移和备份门禁具有明确的运维边界，便于在隔离环境中进行可重复验证并降低生产变更风险。", True, True),
        ("生产部署、生产迁移、真实 Provider 容量和外部告警演练仍需对应环境批准与验收证据。", False, False),
    ]),
    ("前端交互与领域应用", [
        ("前端提供知识库、文档、分块、标签、检索、智能体、会话、设置和管理员等页面，并通过 SSE/WS 展示任务和回答状态。", False, False),
        ("媒体内容通过受控端点访问，浏览器不接收本地路径、模型目录、密钥或未授权任务事件。", False, False),
        ("平台已具备视觉、视频语义段和汽车维修等领域扩展接口；相关完整链路需结合具体模型和受控数据继续验收。", False, False),
        ("项目优势：同一平台既支持通用知识库问答，又能承载视觉、视频和行业应用扩展，便于按专业场景逐步建设能力。", True, True),
    ]),
]

for i, (name, paragraphs) in enumerate(modules, 1):
    add_module(table, i, name, paragraphs)

for row in table.rows:
    for cell in row.cells:
        set_cell_border(cell)

doc.save(OUT)
print(OUT)
