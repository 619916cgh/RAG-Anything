from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path


OUT = Path(__file__).resolve().parents[1] / "多模态知识库与智能体问答平台项目参数清单.docx"


def set_cell_border(cell, color="000000", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)


def margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def font(run, size=10.5, bold=False):
    run.font.name = "宋体"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_width(cell, width):
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_detail(cell, text, *, bold=False, first_indent=False):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(1)
    if first_indent:
        pf.first_line_indent = Pt(18)
    r = p.add_run(text)
    font(r, 10.5, bold=bold)
    return p


rows = [
    ("", "通过 React/Vite 前端、FastAPI 路由层、业务服务、RAG 核心、PostgreSQL 业务状态、向量与图谱存储、Redis 以及文件工作区，形成“知识接入—业务处理—索引构建—混合检索—流式回答—引用回溯”的闭环。"),
    ("知识接入与处理", "1. 支持多种格式接入：文本与办公文档 txt、md、pdf、doc、docx、ppt、pptx、xls、xlsx、html、json；表格与数据 csv、xlsx；图像 jpg、jpeg、png、bmp、gif；音视频 mp3、wav、mp4、webm。\n2. 支持全局解析器选择，并可按文件类型配置覆盖。\n3. 支持 fixed_size、recursive、sentence 等分块策略；按自然边界识别 Markdown 标题章节，保留结构后再按段落细分，并以 token 大小和语义边界控制分块。"),
    ("知识库与内容", "1. 知识库管理功能，支持创建、列表、切换、元数据更新、删除；按能力控制写入、管理和删除。\n2. 文档管理功能，支持列表、摘要分页、下载、删除、批量删除、重试和检索修复。\n3. 分块管理功能，支持分页查看、搜索、编辑、标签更新和删除。\n4. 支持标签列表、文档关联、自动重新生成与恢复；支持实体、节点和边查询与维护；图谱数据可供前端可视化；支持受控媒体交付、检索就绪检查、只读 readiness 预览和文档档案修复。"),
    ("检索与问答", "1. 支持 BM25、向量和图谱三路检索通道，并通过 RRF（RRF_K 默认 60）融合结果。\n2. 支持 balanced、precise、broad 等检索模式，可配置 retrieval_top_k、chunk_top_k、重排和通道裁剪。\n3. Agent SSE 支持 accepted、检索状态、增量 token、引用和完成等事件，会话及关键状态写入 PostgreSQL。\n4. 支持引用回溯、检索健康检查和按知识库范围进行访问控制。"),
    ("智能体", "1. 支持智能体模板、创建、编辑、复制、删除和启用管理。\n2. 支持智能体绑定可访问知识库、配置提示词和检索参数，并按调用者权限校验知识库范围。\n3. 支持会话创建、历史记录、流式回答和共享智能体使用；共享智能体不共享会话历史。\n4. 支持任务 owner 定向事件、脱敏监控通道和连接期间的会话代际复核。"),
    ("用户与安全", "1. 采用 super_admin、dept_admin、teacher、assistant、student 五级 RBAC，权限采用 resource:action 格式。\n2. 所有受保护接口统一经过 require_permission() 和当前持久角色校验。\n3. PostgreSQL 持久化用户、角色、审计、Token 撤销和会话代际信息；WebSocket 收发期间复核撤销状态。\n4. 浏览器不接收本地路径、模型目录、密钥或未授权任务事件。"),
    ("设置与配置", "1. 设置按平台默认、个人设置、知识库资源设置以及请求或任务覆盖进行合并。\n2. 新上传任务固化解析、分块、模型和检索设置快照，历史任务不因后续设置变化而重写。\n3. 设置更新带 revision 和 fingerprint，使用 expected_revision 检测并发冲突，冲突返回 409。"),
    ("部署与运维", "1. 支持 app、postgres、redis、migrate、marker、nginx 等服务组成的 Compose 部署，应用通过内部网络运行。\n2. /api/live 检查进程存活，/api/ready 检查依赖可用性；Prometheus 指标和脱敏监控通道用于运行观察。\n3. PostgreSQL 作为认证、共享状态、迁移历史和业务仓库的权威后端。\n4. 生产迁移前必须完成可恢复备份确认并设置 MIGRATION_BACKUP_ACKNOWLEDGED=true。"),
    ("前端与领域应用", "1. 前端提供知识库、文档、分块、标签、检索、智能体、会话、设置和管理员页面，并通过 SSE/WS 展示任务与回答状态。\n2. 支持图片、音频和视频等受控媒体访问。\n3. 已具备视觉、视频语义段和汽车维修等领域扩展接口，可按专业场景继续扩展。"),
]

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1.25)
section.right_margin = Inches(1.25)

normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
normal.font.size = Pt(10.5)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(6)
r = title.add_run("多模态知识库与智能体问答平台")
font(r, 20)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
grid = table._tbl.tblGrid
for child in list(grid):
    grid.remove(child)
for width in (1.05, 1.05, 4.15):
    col = OxmlElement("w:gridCol")
    col.set(qn("w:w"), str(int(width * 1440)))
    grid.append(col)
header = table.rows[0].cells
for cell, text in zip(header, ("产品名称", "项目", "参数")):
    set_width(cell, 1.05 if cell is not header[2] else 4.15)
    shade(cell, "D9D9D9")
    margins(cell, top=90, bottom=90)
    set_cell_border(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    font(r, 10.5, bold=False)

for project, details in rows:
    cells = table.add_row().cells
    for i, cell in enumerate(cells):
        set_width(cell, (1.05, 1.05, 4.15)[i])
        set_cell_border(cell)
        margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if i < 2 else WD_CELL_VERTICAL_ALIGNMENT.TOP
        if i < 2:
            shade(cell, "D9D9D9")
    p = cells[1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(project)
    font(r, 10.5, bold=True)
    detail = cells[2]
    detail.paragraphs[0]._element.getparent().remove(detail.paragraphs[0]._element)
    for line in details.split("\n"):
        add_detail(detail, line, first_indent=not line[:2].isdigit())

# Merge the product-name column to mirror the source table's vertically merged product cell.
first_body = table.rows[1].cells[0]
for row in table.rows[2:]:
    first_body = first_body.merge(row.cells[0])
first_body.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
shade(first_body, "D9D9D9")
margins(first_body, top=100, bottom=100)
p = first_body.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("多模态知识库与智能体问答平台")
font(r, 10.5, bold=False)
set_cell_border(first_body)

doc.save(OUT)
print(OUT)
