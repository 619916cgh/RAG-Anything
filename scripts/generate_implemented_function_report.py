from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


OUT = Path(__file__).resolve().parents[1] / "RAG-Anything已实现功能参数介绍报告.docx"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9E2F3", size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_font(run, size=10.5, bold=False, color="263238"):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(p, text, size=10.5, bold=False, color="263238"):
    r = p.add_run(text)
    set_font(r, size, bold, color)
    return r


def add_para(doc, text="", style=None, after=5, first_indent=True):
    p = doc.add_paragraph(style=style)
    if text:
        add_text(p, text)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.3
    if first_indent and text and style is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.7 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.2
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        set_cell_shading(cell, "1F4E79")
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, h, 9.2, True, "FFFFFF")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            if ri % 2 == 0:
                set_cell_shading(cell, "F5F8FC")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i else WD_ALIGN_PARAGRAPH.LEFT
            add_text(p, str(value), 9.1, i == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_note(doc, lead, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.right_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    add_text(p, lead, 10, True, "1F4E79")
    add_text(p, text, 10)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(38, 50, 56)
for level, size, color in [(1, 16, "1F4E79"), (2, 12.5, "2F5597"), (3, 11, "3B6EA5")]:
    st = doc.styles[f"Heading {level}"]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(header, "RAG-Anything 功能参数报告", 8.5, False, "6B7C93")
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(footer, "基于当前工作区实现与验证记录编制 · 2026年9月4日", 8, False, "6B7C93")

# Cover
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "RAG-Anything", 29, True, "1F4E79")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "已实现功能参数介绍报告", 22, True, "263238")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, "多模态知识库与智能体问答平台", 12, False, "3B6EA5")
for _ in range(3): doc.add_paragraph()
cover = [
    ("报告性质", "现状说明与参数参考"),
    ("事实基线", "当前代码、迁移、配置与 PROJECT_SUMMARY.md"),
    ("核验日期", "2026年9月4日（Asia/Shanghai）"),
    ("适用读者", "项目管理、技术评审、实施交付与运维人员"),
]
add_table(doc, ["项目项", "当前值"], cover)
add_note(doc, "阅读说明：", "附件 project-brief.docx 是通用英文 Project Brief 占位模板，仅提供 Overview、Goals、Scope、Milestones、Risks and Decisions 的组织结构，不包含本项目事实、参数或验收要求。本报告的功能结论以仓库实际实现为准。")
doc.add_page_break()

# Executive summary
add_heading(doc, "一 执行摘要", 1)
add_para(doc, "RAG-Anything 当前实现是一套面向教育和专业实训场景的多模态知识库与智能体平台。系统通过 React/Vite 前端、FastAPI 路由层、业务服务、RAG 核心、PostgreSQL 业务状态、向量与图谱存储、Redis 以及文件工作区，形成“知识接入—异步处理—索引构建—混合检索—流式回答—引用回溯”的闭环。")
add_para(doc, "本报告描述已经落在当前代码和配置中的能力，不把规划项、历史文档中的旧模型或尚未完成的真实外部验收写成已交付事实。对于代码级测试、隔离 PostgreSQL/Docker 演练、浏览器或真实模型验收，分别列出证据边界。")
add_table(doc, ["能力域", "当前已实现", "状态说明"], [
    ("知识接入与处理", "上传任务、异步 Worker、解析器目录、按任务快照分块与模型设置", "代码与聚焦测试；真实外部模型依赖需单独验收"),
    ("知识库与内容", "KB CRUD、文档/分块/标签、下载删除、视频语义段接口、图谱节点边", "运行时接口已存在；视频真实链路仍按验收门禁单独判断"),
    ("检索与问答", "BM25/向量/图谱通道、RRF 参数、Agent SSE、引用与会话", "SSE 契约和部分隔离演练通过；真实模型路径不由本报告代验"),
    ("智能体", "模板、CRUD、KB 绑定、查询模式、Top-K、重排、会话管理", "参数边界由 Pydantic/路由校验"),
    ("权限与安全", "PostgreSQL 认证、JWT、五级 RBAC、撤销、审计、CORS、限流、请求体限制", "后端权限是最终边界；生产配置仍需受控注入"),
    ("运行与运维", "live/ready 分离、Prometheus 指标、迁移器、Compose app/pg/redis/migrate/marker/nginx", "离线资格与隔离演练不等同生产部署"),
])

# Architecture
add_heading(doc, "二 系统架构与运行链路", 1)
add_heading(doc, "2.1 分层结构", 2)
add_table(doc, ["层次", "职责", "主要实现位置"], [
    ("用户与前端", "知识库、智能体、问答、设置、管理员页面；SSE/WS 客户端", "frontend/src/"),
    ("API 与权限", "REST、WebSocket、SSE 路由；认证、权限依赖、参数校验", "raganything/routers/、dependencies.py"),
    ("业务服务", "KB、上传任务、标签、会话、设置、模型目录、Provider 治理", "raganything/services/"),
    ("RAG 核心", "解析、分块、Embedding、图谱、检索、重排、引用构造", "raganything/parser/、chunking/、query/、processor/"),
    ("持久化与外部组件", "PostgreSQL 权威业务状态；Redis；向量/图谱；文件工作区", "migrations/、pg_* repo、docker-compose.yml"),
    ("异步处理", "上传任务领取、处理、重试、恢复、终态写回", "process_worker.py、services/upload_retry.py"),
])
add_heading(doc, "2.2 关键数据流", 2)
for item in [
    "上传：客户端调用 /api/upload 或 /api/upload/batch，服务端创建带 owner、KB、设置快照和追踪标识的任务。",
    "处理：Worker 读取任务快照，按文件类型选择解析器，执行文本/图像/表格/公式等处理，再进行分块、Embedding、图谱写入与检索就绪状态更新。",
    "查询：请求先经过 JWT/RBAC/KB 能力校验，再按智能体或请求覆盖解析 query_mode、retrieval_top_k、chunk_top_k、enable_rerank 等运行参数。",
    "流式回答：Agent SSE 发送 accepted、检索状态、增量 token、引用/完成等事件；会话和关键状态写入 PostgreSQL，连接期间复核撤销和 session generation。",
    "媒体访问：图片、音频和视频通过受控媒体端点提供，浏览器不接收本地路径、模型目录、密钥或未授权任务事件。",
]: add_bullet(doc, item)

# Ingestion
add_heading(doc, "三 知识接入与文档处理参数", 1)
add_heading(doc, "3.1 支持的接入范围", 2)
add_table(doc, ["类别", "已支持格式或方式", "处理说明"], [
    ("文本与办公文档", "txt、md、pdf、doc、docx、ppt、pptx、xls、xlsx、html、json", "Office 可转 PDF/Markdown；文本可直接读取"),
    ("表格与数据", "csv、xlsx；数据库连接能力由对应服务与配置提供", "保留结构化内容并参与索引"),
    ("图像", "jpg、jpeg、png、bmp、gif 等", "PaddleOCR/VLM 能力按目录与配置启用"),
    ("音视频", "mp3、wav、mp4、webm 等媒体类型", "媒体交付与视频语义段接口已存在；真实模型链路需独立验收"),
    ("扩展解析器", "MinerU、Docling、Marker、PaddleOCR、OpenDataLoader；支持注册自定义 Parser", "parser/__init__.py 提供工厂与注册机制"),
])
add_heading(doc, "3.2 解析器与处理开关", 2)
add_table(doc, ["参数", "默认或取值", "作用"], [
    ("PARSER", "docling", "全局解析器选择；可按文件类型配置覆盖"),
    ("ENABLE_IMAGE_PROCESSING", "true", "是否处理图像内容"),
    ("ENABLE_TABLE_PROCESSING", "true", "是否处理表格内容"),
    ("ENABLE_EQUATION_PROCESSING", "true", "是否处理公式内容"),
    ("ENTITY_TYPES", "空（按实现默认）", "实体抽取类型配置"),
    ("ENTITY_EXTRACTION_MIN_DEGREE", "0", "实体关系最小度阈值"),
    ("MAX_UPLOAD_SIZE_MB", "500", "上传请求体上限；非上传请求由 MAX_BODY_SIZE_MB 控制"),
    ("PROCESS_TIMEOUT", "3600 秒（历史兼容入口）", "任务处理时限；Worker 另有空闲超时与租约控制"),
])
add_note(doc, "重要边界：", "仓库安全核验指出上传 API 尚未统一扩展名白名单，对外错误也不是全局统一脱敏。因此，本报告不把“所有上传均受扩展名白名单保护”写成已实现能力。")
add_heading(doc, "3.3 分块策略", 2)
add_table(doc, ["策略", "核心行为", "默认参数或成本"], [
    ("fixed_size", "LightRAG 固定大小切割与重叠窗口", "作为默认回退；无额外外部调用"),
    ("recursive", "段落→换行→句子→字符的自然边界递归分割", "CHUNK_SIZE=800；overlap 参数默认 100 token（运行设置默认 chunk_size=800）"),
    ("sentence", "按中英文句末边界拆句，再合并到目标 token 大小", "目标 800；长句按 token 窗口切分"),
    ("structure", "识别 Markdown 标题/章节，保留结构后再按段落细分", "目标 800；适合章节清晰文档"),
    ("semantic", "段落 Embedding 相似度检测语义边界", "Embedding 批次最多 10 条；失败回退 recursive"),
    ("agentic", "LLM 分析话题转换并给出分割点", "最多分析 500 行；单次 max_tokens=500；失败回退 recursive"),
])
add_note(doc, "设置快照：", "新上传任务会固化平台、个人、KB 和单次覆盖合并后的解析/分块/模型/检索设置；历史任务快照不因后续设置变化而重写。")

# KB and retrieval
add_heading(doc, "四 知识库、内容管理与检索参数", 1)
add_heading(doc, "4.1 知识库与内容生命周期", 2)
add_table(doc, ["功能", "已实现行为", "关键接口"], [
    ("知识库管理", "创建、列表、切换、元数据更新、删除；按能力控制写入/管理/删除", "/api/kb/list、/api/kb/create、/api/kb/{kb}/metadata"),
    ("文档管理", "列表、摘要分页、下载、删除、批量删除、重试、检索修复", "/api/knowledge/documents*"),
    ("分块管理", "分页查看、搜索、编辑、标签更新、删除", "/api/knowledge/documents/{id}/chunks*"),
    ("标签", "标签列表、文档关联、自动重新生成与恢复", "/api/knowledge/tags*、.../tags/regenerate"),
    ("图谱", "实体/节点/边查询与增删改；图谱数据供前端可视化", "/api/knowledge/graph*"),
    ("媒体", "受控图片/音视频交付；视频语义段查询端点", "/api/knowledge/media*、.../video-segments"),
    ("健康与修复", "检索就绪检查、只读 readiness 预览、文档检索修复", "/api/knowledge/retrieval-health、.../repair-retrieval"),
])
add_heading(doc, "4.2 混合检索与 RRF", 2)
add_table(doc, ["参数", "默认值", "说明"], [
    ("preset", "balanced", "可选 balanced、precise、broad，解析为固定可审计参数"),
    ("channels", "bm25、vector、graph", "三路检索通道，可按设置策略裁剪"),
    ("RRF_K", "60", "Reciprocal Rank Fusion 的常数"),
    ("BM25_TOP_K", "50", "关键词通道候选数"),
    ("VECTOR_TOP_K", "100", "向量通道候选数"),
    ("GRAPH_TOP_K", "30", "图谱通道候选数"),
    ("GRAPH_DEPTH", "2", "图谱邻居遍历深度；平台限制范围 0–64"),
    ("BM25_K1 / BM25_B", "1.5 / 0.75", "BM25 词频与长度归一化参数"),
    ("BM25_TOKENIZER", "jieba", "中文分词器配置"),
    ("RRF_CHANNEL_TIMEOUT", "0.15 秒", "单通道响应预算"),
])
add_heading(doc, "4.3 向量、图谱与重排", 2)
add_table(doc, ["能力", "当前实现参数"], [
    ("Embedding", "默认 EMBEDDING_MODEL=text-embedding-v3，EMBEDDING_DIM=1024；实际模型目录由服务端 catalog 决定"),
    ("向量相似度", "余弦相似度；PostgreSQL HNSW 部署参数由 POSTGRES_HNSW_M/POSTGRES_HNSW_EF 控制，Compose 默认 16/64"),
    ("知识图谱", "实体、关系、属性和关联 chunk；支持实体查询、节点/边维护与前端图谱数据接口"),
    ("Rerank", "请求或智能体 enable_rerank 开关；入口使用服务端配置的 rerank 模型与候选预算"),
    ("引用", "回答管线保留来源文档、分段和媒体引用；前端可根据引用跳转受控资源"),
])

# Agents
add_heading(doc, "五 智能体与问答功能参数", 1)
add_heading(doc, "5.1 智能体配置模型", 2)
add_table(doc, ["字段", "默认值或范围", "作用"], [
    ("name / icon / description", "名称默认“新智能体”；icon 默认 🤖", "展示与识别信息"),
    ("kb_name", "default", "绑定知识库"),
    ("llm_model", "qwen-plus", "LLM 模型标识；生产推荐通过模型 profile 目录选择"),
    ("temperature", "0.0", "生成随机性，由模型提供方约束"),
    ("max_response_tokens", "4096", "回答最大 token 数"),
    ("query_mode", "hybrid", "hybrid、local、global、naive 等检索模式；Agentic 内部可映射 rrf"),
    ("agent_mode", "none", "普通问答或 Agentic 查询模式"),
    ("retrieval_top_k", "40，路由限制 5–200", "检索候选规模"),
    ("chunk_top_k", "20", "上下文分段规模"),
    ("enable_rerank", "false", "是否启用重排"),
    ("include_references", "true", "是否返回引用信息"),
    ("system_prompt / use_default_prompt", "空 / true", "系统提示词与默认提示词开关"),
])
add_heading(doc, "5.2 预置模板", 2)
add_table(doc, ["模板", "模式", "Top-K / 温度", "适用方向"], [
    ("通用文档问答", "hybrid", "40 / 0.0", "通用知识问答"),
    ("学术论文助手", "hybrid", "60 / 0.1", "严谨引用与文献分析"),
    ("代码文档助手", "hybrid", "40 / 0.1", "代码示例与技术细节"),
    ("精确检索器", "local", "20 / 0.0", "仅返回直接相关信息"),
    ("图谱实体推理", "graph", "30 / 0.0", "实体关系路径追溯"),
    ("创意头脑风暴", "global", "80 / 0.7", "跨领域联想与方案发散"),
])
add_heading(doc, "5.3 会话与流式协议", 2)
add_table(doc, ["项目", "当前协议"], [
    ("会话", "智能体可创建、查询、更新、删除对话线程和消息；共享智能体不共享会话历史"),
    ("SSE", "POST /api/agents/{agent_id}/query/stream；事件包含 accepted、heartbeat/状态、retrieval、增量 token、done 或错误"),
    ("超时", "前端客户端具备首事件超时与空闲超时；Nginx 配置无缓冲转发；服务端在会话期间复核 token 撤销与 session generation"),
    ("WebSocket", "保留 /api/ws、/api/ws/monitor 等受控通道；失效会话以 4001 关闭，监控通道仅发送聚合脱敏字段"),
])

# RBAC security
add_heading(doc, "六 用户、权限与安全参数", 1)
add_heading(doc, "6.1 五级 RBAC", 2)
add_table(doc, ["角色", "核心权限范围", "管理边界"], [
    ("super_admin", "全部 Permission", "平台、用户、KB、智能体、设置、审计与监控"),
    ("dept_admin", "用户读写、KB 管理、智能体、设置读、审计、监控、图谱、AutoRepair", "系部范围内的管理操作"),
    ("teacher", "KB 读写管理、智能体读写、监控/分析、图谱、AutoRepair", "自有业务管理能力由后端能力控制"),
    ("assistant", "KB 读写、智能体读取、监控、图谱、AutoRepair 读取", "内容维护与使用"),
    ("student", "KB 读取、智能体读取、图谱读取、AutoRepair 读取", "只读与问答"),
])
add_note(doc, "权限原则：", "前端隐藏只是体验层，所有受保护接口统一经过 require_permission() 和当前持久角色校验。当前运行时已退役旧 allowed_kbs 成员授权模型，KB 可见性与写入/管理/删除能力按全局角色权限和能力投影执行。")
add_heading(doc, "6.2 认证与会话", 2)
add_table(doc, ["参数", "当前实现"], [
    ("存储", "PostgreSQL 用户、角色、审计、Token 撤销与 session generation 为权威状态"),
    ("密码", "bcrypt 哈希；密码策略由认证服务校验"),
    ("JWT", "HS256；Access Token 默认 JWT_EXPIRY_HOURS=1 小时；Refresh Token 默认 REFRESH_EXPIRY_DAYS=7 天"),
    ("刷新", "Refresh Token 家族轮转与重放检测；旧 token 并发消费失败时撤销 family"),
    ("暴力破解", "MAX_FAILED_LOGIN_ATTEMPTS=5；LOGIN_LOCKOUT_MINUTES=15"),
    ("注销/改密", "撤销 token 并提升 session generation，使已建立的 SSE/WS 会话停止继续发送受保护事件"),
])
add_heading(doc, "6.3 网络与请求安全", 2)
add_table(doc, ["项目", "当前值或行为"], [
    ("CORS", "显式来源白名单；允许 GET/POST/PUT/PATCH/DELETE/OPTIONS；禁止通配 *"),
    ("限流", "默认全局 120/minute；登录与刷新等敏感端点采用更严格独立限制"),
    ("请求体", "非上传默认 MAX_BODY_SIZE_MB=10；上传默认 MAX_UPLOAD_SIZE_MB=500"),
    ("响应头", "nosniff、DENY、HSTS、Referrer-Policy、Permissions-Policy、CSP 等安全头"),
    ("日志", "apply_sensitive_log_filter / mask_sensitive_data 对 password、token、secret、key、api_key 等脱敏"),
    ("审计", "权限拒绝、用户/平台设置更新等关键操作写入审计；支持管理员查询与健康审计端点"),
])

# Settings
add_heading(doc, "七 个人与平台设置", 1)
add_para(doc, "设置服务使用 PostgreSQL 持久化，并以平台默认、旧环境兼容层、资源级设置、个人设置和请求/任务覆盖的优先级合并。设置返回带 revision 和 fingerprint，更新使用 expected_revision，冲突返回 409，避免并发覆盖。")
add_table(doc, ["设置分区", "可配置内容", "权限"], [
    ("models", "LLM/VLM profile；profile 由服务端目录公开投影，凭据与传输参数不下发", "agent:write"),
    ("ingestion", "parser、chunking_strategy、chunk_size、图像/表格/公式开关、按类型 parser 覆盖", "kb:write"),
    ("retrieval", "balanced/precise/broad 预设、BM25/向量/图谱 Top-K、深度、通道与分词器", "kb:write"),
    ("runtime", "LLM timeout、个人并发等运行配额", "kb:write"),
    ("platform", "平台默认、允许目录、限制范围、只读状态与 retrieval preset 版本", "settings:read/write"),
])
add_table(doc, ["平台限制", "允许范围"], [
    ("worker_concurrency / provider_concurrency / personal_concurrency", "1–1024"),
    ("llm_timeout", "1–3600 秒"),
    ("bm25_top_k / vector_top_k / graph_top_k", "1–10000"),
    ("graph_depth", "0–64"),
    ("cache_capacity", "1–1,000,000"),
    ("interactive_wait_seconds", "0–300 秒"),
])

# API
add_heading(doc, "八 API 与前端功能面", 1)
add_para(doc, "当前路由按领域拆分为 auth、admin、knowledge、agent、demo、autorepair、vision 和 user_settings。以下为面向实施与集成最有价值的接口族，不把已移除的 DAG 路由列入当前能力。")
add_table(doc, ["接口族", "代表端点", "用途"], [
    ("认证与用户", "/api/auth/register、login、refresh、me、logout；/api/admin/users、roles、audit-logs", "登录、会话轮转、用户角色和审计"),
    ("知识库与上传", "/api/kb/list、create、switch；/api/upload、upload/batch、upload/tasks", "KB 生命周期、异步上传任务与重试取消"),
    ("文档与分块", "/api/knowledge/documents、document-summaries、{id}/chunks、download、batch-delete", "文档列表、分段管理和源文件访问"),
    ("检索与图谱", "/api/knowledge/retrieval-health、graph、entities、image-search", "检索状态、实体关系与多模态搜索"),
    ("智能体", "/api/agents、templates、{id}/conversations、{id}/query/stream", "模板、CRUD、会话和流式问答"),
    ("设置与模型", "/api/users/me/settings*、/api/admin/platform、/api/model-profiles、/api/vision-models", "个人/平台设置与模型目录"),
    ("领域应用", "/api/autorepair/*、/api/demo/*", "AutoRepair 知识图谱/问答与受控公开演示"),
    ("健康与监控", "/api/live、/api/ready、/api/health、/api/monitor/*、/api/ws/monitor", "进程存活、依赖就绪、指标与监控"),
])

# Deployment
add_heading(doc, "九 部署、存储与运维参数", 1)
add_heading(doc, "9.1 Compose 服务", 2)
add_table(doc, ["服务", "职责", "关键配置"], [
    ("app", "FastAPI 应用与业务服务", "内部监听 8000；依赖 migrate、postgres、redis、marker"),
    ("postgres", "业务、认证、迁移历史与共享状态", "postgres:16-alpine；healthcheck 10s/5 retries"),
    ("redis", "缓存与部分会话/队列辅助", "redis:7-alpine；持久卷 redisdata"),
    ("migrate", "启动前执行 PostgreSQL 迁移", "必须 MIGRATION_BACKUP_ACKNOWLEDGED=true；advisory lock 单会话执行"),
    ("marker", "隔离的文档解析服务", "健康检查 30s；默认 8g/2 CPU，可配置"),
    ("nginx", "静态前端、反向代理、外部入口", "宿主绑定 NGINX_BIND_PORT，app 不直接暴露宿主端口"),
])
add_heading(doc, "9.2 健康检查与运行指标", 2)
add_table(doc, ["端点/指标", "语义"], [
    ("/api/live", "只检查进程存活，不把短暂 PostgreSQL 瞬断等同于进程死亡"),
    ("/api/ready", "检查依赖与应用是否可接收流量"),
    ("/api/health", "综合健康信息，供应用和外部探针使用"),
    ("Prometheus", "存储字节数、文件数、运行时与 Provider 治理指标；scrape inventory 已纳入仓库"),
    ("日志", "RotatingFileHandler 默认 10MB、保留 10 个备份；LOG_LEVEL 可配置"),
])
add_heading(doc, "9.3 迁移与备份门禁", 2)
add_para(doc, "迁移链当前包含 001–037 的 PostgreSQL 演进文件，包含 RBAC、Token 撤销、上传重试、设置、视频语义段、模型身份、公开演示、Provider 治理和 DAG 移除等历史变更。生产 apply 前必须先有可恢复备份，并显式设置 MIGRATION_BACKUP_ACKNOWLEDGED=true。037_remove_dag_workflow.sql 会删除历史 DAG 定义与运行表，属于不可逆数据变更，不能以普通功能开关替代备份与审批。")

# Verification
add_heading(doc, "十 验证证据与未覆盖边界", 1)
add_heading(doc, "10.1 可引用的验证层级", 2)
add_table(doc, ["证据层级", "本项目已有事实", "不能推出的结论"], [
    ("代码与聚焦测试", "权限、设置、SSE 契约、迁移器、Provider 治理、前端工具等有针对性测试记录", "不能推出真实外部模型或生产行为"),
    ("构建与静态检查", "Python 编译、前端 locked build、OpenSpec strict、summary/diff checks 有记录", "不能推出浏览器完整 UAT 或部署健康"),
    ("隔离 PostgreSQL/Docker", "Stage 3 隔离 rehearsal 记录包含迁移、Worker、检索、SSE/citation、五级 RBAC、恢复与 rollback", "run-owned 资源成功不等于线上资源已部署"),
    ("真实 Provider", "真实 Qwen 阶梯曾运行到 50 VU，但 50 阶段 TTFT P95 超过 5 秒门禁并硬停止", "不能宣称 50 VU 或生产容量通过"),
    ("浏览器/Worker/生产", "部分场景有历史或隔离证据，当前 active change 仍需按门禁复核", "不能将本报告作为当前生产验收签字"),
])
add_heading(doc, "10.2 当前明确的限制与风险", 2)
for item in [
    "外部真实 LLM、Embedding、VLM、ASR 的可达性、配额与延迟受部署环境影响，需用受控账号和真实 Worker 单独验收。",
    "视频语义段代码与接口存在，但最近的正式浏览器/视频诊断曾在资源契约或 citation 渲染前停止，不能宣称完整 Video 5.4 端到端通过。",
    "Prometheus staging scrape、Alertmanager 演练、Docker PostgreSQL 抖动、Windows Worker 后代进程等仍是外部或受控环境任务。",
    "当前工作区存在并行任务改动；本报告仅读取并引用事实，不将这些未提交 harness 改动包装成产品功能。",
    "旧资料中的 SQLite 认证、三角色模型、旧 /upload 与 /query 描述、DAG 工作流不代表当前实现。",
]: add_bullet(doc, item)

# Vision, domain application and frontend
add_heading(doc, "十一 视觉视频、领域应用与前端交互", 1)
add_heading(doc, "11.1 视觉与视频能力", 2)
add_table(doc, ["能力", "已实现参数或行为", "验收边界"], [
    ("视觉模型目录", "服务端维护 VLM、Embedding profile；公开 API 只返回无密钥投影", "真实模型可用性需按部署目录与凭据单独探针"),
    ("视频片段", "默认目标片段 24 秒，范围 15–30 秒，重叠 3 秒；结合 ASR/场景边界生成确定性片段 ID", "代码与接口存在；完整浏览器 Video Phase A 仍需复核"),
    ("视频资源", "默认最大时长 3600 秒、最多 60 帧；视频/帧/片段并发由配置控制，媒体通过受控端点访问", "不把未完成的真实 Worker/Provider 链路写成生产验收"),
    ("图片搜索", "提供 /api/image-search 与媒体访问端点，可返回受控引用和相似结果", "需要视觉 Embedding profile 可用"),
])
add_heading(doc, "11.2 AutoRepair 汽修领域应用", 2)
add_table(doc, ["模块", "功能参数"], [
    ("知识图谱", "图谱摘要、节点/边浏览、实体谱系与关联路径；前端可按领域 KB 查询"),
    ("案例库", "故障案例、工艺案例统一搜索、分类、统计与 CRUD；案例搜索默认 Top-K 20"),
    ("代码解析", "支持 G 代码/PLC 指令解析接口，返回结构化指令与诊断信息"),
    ("领域问答", "普通问答与 SSE 流式问答；默认最多 3 步的受控 QA/诊断流程"),
    ("故障诊断", "支持多轮继续诊断接口；权限按 autorepair:read/write 控制"),
])
add_heading(doc, "11.3 前端交互参数", 2)
add_table(doc, ["交互项", "当前值或行为"], [
    ("通用分页", "常用页大小 10/20/50；分块页面 25/50/100，默认 25"),
    ("上传请求", "客户端超时 600 秒；任务状态通过轮询与 SSE/WS 反馈"),
    ("SSE 客户端", "首事件超时 20 秒，空闲超时 45 秒；支持 accepted/heartbeat/done 边界"),
    ("缓存", "KB 列表缓存 5 秒；详情/统计缓存 30 秒，详情缓存最多 20 项"),
    ("监控与任务", "监控页约 5 秒刷新；KB 任务约 15 秒轮询；取消任务约 2 秒轮询"),
    ("界面范围", "知识库、文档详情/分块、智能体、AutoRepair、个人设置、平台设置、用户、审计、监控、公开演示"),
])

# Appendix
add_heading(doc, "附录 A 事实来源与建议核对路径", 1)
add_table(doc, ["来源", "用于核对的内容"], [
    ("PROJECT_SUMMARY.md", "项目定位、当前能力、架构、业务规则、验证边界、近期任务"),
    ("server.py", "FastAPI 组装、默认模型/Embedding、CORS、限流、请求体、安全头、日志与健康端点"),
    ("raganything/routers/", "API 路由、Pydantic 参数、权限依赖、SSE/WS 端点"),
    ("raganything/services/user_settings.py", "四类设置、预设、平台限制、revision/fingerprint 和快照合并"),
    ("raganything/services/pg_auth_repo.py 与 permissions.py", "PostgreSQL 认证、JWT、五级 RBAC、锁定与会话生命周期"),
    ("raganything/chunking/__init__.py", "六类分块策略及回退、批次和分析上限"),
    ("raganything/parser/", "解析器注册表、内置解析器和自定义 Parser 扩展点"),
    ("docker-compose.yml 与 migrations/", "服务编排、健康检查、卷、迁移门禁与历史演进"),
    ("tests/ 与 PROJECT_SUMMARY 近期记录", "代码级、隔离验收、浏览器/Provider 证据边界"),
])
add_note(doc, "交付结论：", "本报告可作为当前版本的功能参数说明和实施沟通基线。涉及生产部署、生产迁移、真实模型容量、Worker/浏览器完整链路的判断，应在对应受控环境证据齐全后另行出具验收结论。")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
